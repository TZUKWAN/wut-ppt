# -*- coding: utf-8 -*-
"""
build_deck.py — 武汉理工行政汇报 PPT 生成器（高密度逻辑图版）

用法: python scripts/build_deck.py 大纲.json 输出.pptx

设计规约（用户五条硬要求）：
1. 全文最小字号 ≥ 12pt
2. 每个要点必须用独立形状包裹（卡片+序号圆徽），不允许纯文字列表
3. 必须有视觉逻辑引导（箭头链/连接线/时间轴），形状多样：
   圆、圆角矩形、箭头(CHEVRON/RIGHT_ARROW)、六边形、五边形、平行四边形
4. 高信息密度：每页要点带完整解释句
5. 不允许大片空白：zone 垂直流自动填满正文区

文字标记：**加粗蓝**  [[加粗红]]
"""
import json, math, os, re, sys
from copy import deepcopy

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(BASE, "assets", "template.pptx")

# ---------------- 设计常量 ----------------
BLUE  = RGBColor(0x00, 0x46, 0x9A)
BLUE_D= RGBColor(0x00, 0x33, 0x71)
BLUE_L= RGBColor(0x3D, 0x6F, 0xB0)
GOLD  = RGBColor(0xFB, 0xC5, 0x40)
GOLD_D= RGBColor(0x9A, 0x73, 0x00)
RED   = RGBColor(0xC0, 0x00, 0x00)
INK   = RGBColor(0x26, 0x26, 0x26)
GRAY  = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD  = RGBColor(0xF4, 0xF7, 0xFB)
LINE  = RGBColor(0xD9, 0xD9, 0xD9)
CHIP  = RGBColor(0xEA, 0xF1, 0xF9)

FONT = "微软雅黑"

PAGE_L, PAGE_R = 0.45, 12.88
PAGE_W = PAGE_R - PAGE_L
BODY_T, BODY_B = 1.02, 6.95

CH_NUM = "一二三四五"

# ---------------- 基础工具 ----------------

def _set_run(run, text, size, bold=False, color=INK):
    run.text = text
    f = run.font
    f.size = Pt(size); f.bold = bold; f.name = FONT
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", FONT)

_MARK = re.compile(r"(\*\*.+?\*\*|\[\[.+?\]\])")

def _segments(text):
    """**x**=蓝粗  [[x]]=红粗  普通段颜色 None（用默认色）"""
    out = []
    for part in _MARK.split(str(text)):
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, BLUE))
        elif part.startswith("[[") and part.endswith("]]"):
            out.append((part[2:-2], True, RED))
        else:
            out.append((part, False, None))
    return out

def _plain(text):
    return re.sub(r"\*\*|\[\[|\]\]", "", str(text))

def add_par(tf, text, size=12, align=PP_ALIGN.LEFT, first=False,
            space_after=3, line=1.22, color=None, bold=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    for seg_text, seg_bold, seg_color in _segments(text):
        r = p.add_run()
        _set_run(r, seg_text, size, bold or seg_bold,
                 seg_color if seg_color is not None else (color or INK))
    return p

def shape(slide, mso, l, t, w, h, fill=None, line_color=None, lw=0.75, adj=None):
    s = slide.shapes.add_shape(mso, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line_color is None: s.line.fill.background()
    else: s.line.color.rgb = line_color; s.line.width = Pt(lw)
    s.shadow.inherit = False
    if adj is not None:
        try: s.adjustments[0] = adj
        except Exception: pass
    return s

def box(slide, l, t, w, h, fill=None, line_color=None, lw=0.75):
    return shape(slide, MSO_SHAPE.RECTANGLE, l, t, w, h, fill, line_color, lw)

def rbox(slide, l, t, w, h, fill=None, line_color=None, lw=0.75, radius=0.055):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill,
                 line_color, lw, adj=radius)

def badge(slide, l, t, d, text, fill=BLUE, tcolor=WHITE, size=12.5):
    """序号圆徽"""
    b = shape(slide, MSO_SHAPE.OVAL, l, t, d, d, fill=fill)
    tf = b.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); _set_run(r, text, size, True, tcolor)
    return b

def text_box(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    return tf

def shape_text(shp, paras, size=12, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
               m=(0.10, 0.08, 0.05, 0.05), color=None, bold=False):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_right = Inches(m[0]), Inches(m[1])
    tf.margin_top, tf.margin_bottom = Inches(m[2]), Inches(m[3])
    lines = paras if isinstance(paras, list) else [paras]
    for i, ln in enumerate(lines):
        add_par(tf, ln, size=size, align=align, first=(i == 0), space_after=2,
                color=color, bold=bold)

# ---------------- 高度估算 ----------------

def _disp_len(text):
    return sum(1.0 if ord(c) > 0x2E7F else 0.55 for c in _plain(text))

def est_lines(text, width_in, pt):
    cpl = max(4.0, width_in * 72.0 / pt)
    return max(1, math.ceil(_disp_len(text) / cpl))

def _lh(pt, line=1.24):
    return pt * line / 72.0

def _split_item(it):
    """要点 → (引导语, 正文)。dict 或 '引导：正文' 或纯文本"""
    if isinstance(it, dict):
        return it.get("lead", ""), it.get("text", "")
    s = str(it)
    if "：" in s:
        lead, text = s.split("：", 1)
        if len(_plain(lead)) <= 12:
            return lead, text
    return "", s

def est_zone(z, w=PAGE_W):
    zt = z["type"]
    if zt == "lead":
        return est_lines(z["text"], w - 0.3, 13) * _lh(13) + 0.20
    if zt == "kv_rows":
        return sum(max(0.36, est_lines(r["text"], w - 1.85, 12.5) * _lh(12.5) + 0.16)
                   for r in z["rows"])
    if zt == "cards":
        n = len(z["cards"]); cw = (w - (n - 1) * 0.16) / n
        body = 0
        for c in z["cards"]:
            h = 0
            for it in c["items"]:
                lead, text = _split_item(it)
                h += est_lines((lead + "：" if lead else "") + text, cw - 0.62, 12) * _lh(12) + 0.18
            body = max(body, h + (len(c["items"]) - 1) * 0.10)
        return 0.42 + body + 0.10
    if zt == "flow_chain":
        n = len(z["items"])
        pitch = (w + (n - 1) * 0.28) / n
        if not any(it.get("text") for it in z["items"]):
            return 0.62
        body = max(est_lines(it["text"], pitch - 0.62, 12) for it in z["items"])
        return 0.55 + 0.14 + body * _lh(12) + 0.22
    if zt == "shape_row":
        n = len(z["items"])
        sw = min(2.0, (w - (n - 1) * 0.3) / n)
        txt = max((est_lines(it.get("text", ""), sw + 0.5, 12) for it in z["items"]), default=1)
        h = 0.72 + (txt * _lh(12) + 0.06 if any(it.get("text") for it in z["items"]) else 0)
        if z.get("label"): h += 0.38
        return h
    if zt == "badge_grid":
        cols = z.get("cols", 3)
        rows = math.ceil(len(z["items"]) / cols)
        cw = (w - (cols - 1) * 0.16) / cols
        cell = max(est_lines(it["text"], cw - 0.24, 12) for it in z["items"])
        return rows * (0.34 + cell * _lh(12) + 0.20) + (rows - 1) * 0.14
    if zt == "compare_rows":
        h = 0.32
        mw = 5.3 - 0.2; rw = w - 1.7 - 5.6 - 0.5 - 0.2
        for r in z["rows"]:
            h += max(est_lines(r["mid"], mw, 12), est_lines(r["right"], rw, 12)) * _lh(12) + 0.20
        return h + (len(z["rows"]) - 1) * 0.12
    if zt == "timeline":
        n = len(z["phases"]); cw = (w - (n - 1) * 0.25) / n
        body = 0
        for ph in z["phases"]:
            h = sum(est_lines(tk, cw - 0.62, 12) * _lh(12) + 0.16 for tk in ph["tasks"])
            body = max(body, h + (len(ph["tasks"]) - 1) * 0.10)
        return 1.06 + body + 0.10
    if zt == "chips":
        rows, x = 1, 0.0
        for it in z["items"]:
            iw = _disp_len(it) * 0.165 + 0.78
            if x + iw > w: rows += 1; x = 0
            x += iw + 0.14
        h = rows * 0.40 + 0.08
        if z.get("label"): h += 0.36
        if z.get("text"): h += est_lines(z["text"], w, 12) * _lh(12) + 0.10
        return h
    if zt == "table":
        widths = z.get("widths") or [1] * len(z["head"])
        tot = float(sum(widths)); h = 0.42
        for row in z["rows"]:
            rh = max(est_lines(v, w * widths[ci] / tot - 0.2, 12) for ci, v in enumerate(row))
            h += rh * _lh(12, 1.25) + 0.16
        return h
    if zt == "banner":
        return 0.46 + len(z.get("lines", [])) * 0.32 + 0.16
    if zt == "note":
        return max(0.46, est_lines(z["text"], w - 0.5, 12.5) * _lh(12.5) + 0.22)
    return 0.4

# ---------------- Zone 渲染 ----------------

def render_lead(slide, z, l, t, w, h):
    box(slide, l, t + 0.04, 0.055, h - 0.08, fill=BLUE)
    tf = text_box(slide, l + 0.18, t, w - 0.18, h, MSO_ANCHOR.MIDDLE)
    add_par(tf, z["text"], size=13, first=True, line=1.32)

def render_kv_rows(slide, z, l, t, w, h):
    rows = z["rows"]; gap = 0.10
    rh = (h - (len(rows) - 1) * gap) / len(rows)
    y = t
    for row in rows:
        chip = box(slide, l, y, 1.55, rh, fill=BLUE)
        shape_text(chip, row["label"], size=12.5, align=PP_ALIGN.CENTER,
                   m=(0.04, 0.04, 0.02, 0.02), color=WHITE, bold=True)
        body = box(slide, l + 1.65, y, w - 1.65, rh, fill=CARD)
        shape_text(body, row["text"], size=12.5, anchor=MSO_ANCHOR.MIDDLE,
                   m=(0.12, 0.10, 0.04, 0.04))
        y += rh + gap

def render_cards(slide, z, l, t, w, h):
    """列卡片：列头(圆角蓝条) + 每个要点一个独立圆角卡 + 序号圆徽，填满整列"""
    cards = z["cards"]; n = len(cards); gap = 0.16
    cw = (w - (n - 1) * gap) / n
    head_h = 0.40
    for i, c in enumerate(cards):
        x = l + i * (cw + gap)
        hd = rbox(slide, x, t, cw, head_h, fill=BLUE, radius=0.14)
        shape_text(hd, c["head"], size=13, align=PP_ALIGN.CENTER,
                   m=(0.06, 0.06, 0.02, 0.02), color=WHITE, bold=True)
        items = c["items"]
        ests = []
        for it in items:
            lead, text = _split_item(it)
            ests.append(est_lines((lead + "：" if lead else "") + text, cw - 0.62, 12) * _lh(12) + 0.18)
        area_t, area_b = t + head_h + 0.12, t + h
        igap = 0.10
        total = sum(ests) + igap * (len(items) - 1)
        extra = max(0.0, (area_b - area_t) - total)
        y = area_t
        for j, it in enumerate(items):
            ih = ests[j] + extra / len(items)
            card = rbox(slide, x, y, cw, ih, fill=CARD, line_color=LINE, lw=0.75)
            d = 0.32
            badge(slide, x + 0.10, y + (ih - d) / 2, d, str(j + 1))
            tf = text_box(slide, x + 0.52, y + 0.03, cw - 0.62, ih - 0.06,
                          MSO_ANCHOR.MIDDLE)
            lead, text = _split_item(it)
            content = (f"**{lead}**：{text}" if lead else text)
            add_par(tf, content, size=12, first=True, line=1.26)
            y += ih + igap

def render_flow_chain(slide, z, l, t, w, h):
    """箭头链：CHEVRON 依次咬合；带 text 时下方对齐详情卡"""
    items = z["items"]; n = len(items)
    has_detail = any(it.get("text") for it in items)
    chev_h = 0.55 if has_detail else min(0.62, h)
    pitch = (w + (n - 1) * 0.28) / n
    palette = [BLUE, BLUE, BLUE, BLUE, RED] if n == 5 else [BLUE] * n
    if z.get("colors"): palette = [globals()[c] for c in z["colors"]]
    for i, it in enumerate(items):
        x = l + i * (pitch - 0.28)
        fill = palette[i] if i < len(palette) else BLUE
        cv = shape(slide, MSO_SHAPE.CHEVRON, x, t, pitch, chev_h, fill=fill)
        tcol = WHITE if fill != GOLD else BLUE_D
        shape_text(cv, it["label"], size=12.5, align=PP_ALIGN.CENTER,
                   m=(0.30, 0.14, 0.02, 0.02), color=tcol, bold=True)
    if has_detail:
        ct = t + chev_h + 0.14
        ch = h - chev_h - 0.14
        for i, it in enumerate(items):
            cx = l + i * (pitch - 0.28) + 0.06
            cwd = pitch - 0.40
            card = rbox(slide, cx, ct, cwd, ch, fill=CARD, line_color=LINE, lw=0.75)
            box(slide, cx, ct, 0.05, ch, fill=palette[i] if i < len(palette) else BLUE)
            tf = text_box(slide, cx + 0.14, ct + 0.06, cwd - 0.22, ch - 0.12,
                          MSO_ANCHOR.MIDDLE)
            add_par(tf, it.get("text", ""), size=12, first=True, line=1.26)

def render_shape_row(slide, z, l, t, w, h):
    """形状横排+连接线：六边形/圆形/五边形，name 在内、text 在下"""
    items = z["items"]; n = len(items)
    y = t
    if z.get("label"):
        box(slide, l, y + 0.03, 0.055, 0.26, fill=GOLD)
        tf = text_box(slide, l + 0.14, y, w - 0.14, 0.34, MSO_ANCHOR.MIDDLE)
        add_par(tf, z["label"], size=13, first=True, color=BLUE_D, bold=True)
        y += 0.38
    mso = {"hexagon": MSO_SHAPE.HEXAGON, "oval": MSO_SHAPE.OVAL,
           "pentagon": MSO_SHAPE.PENTAGON, "chevron": MSO_SHAPE.CHEVRON,
           "parallelogram": MSO_SHAPE.PARALLELOGRAM}.get(z.get("shape", "hexagon"),
                                                         MSO_SHAPE.HEXAGON)
    gap = 0.30
    sw = min(2.0, (w - (n - 1) * gap) / n)
    sh = 0.72
    total_w = n * sw + (n - 1) * gap
    x0 = l + (w - total_w) / 2
    # 连接线
    box(slide, x0 + sw / 2, y + sh / 2 - 0.012, total_w - sw, 0.024, fill=BLUE_L)
    for i, it in enumerate(items):
        x = x0 + i * (sw + gap)
        sp = shape(slide, mso, x, y, sw, sh, fill=BLUE)
        shape_text(sp, it["name"], size=12.5, align=PP_ALIGN.CENTER,
                   m=(0.16, 0.16, 0.02, 0.02), color=WHITE, bold=True)
        if it.get("text"):
            tf = text_box(slide, x - 0.25, y + sh + 0.05, sw + 0.5, 0.55,
                          MSO_ANCHOR.TOP)
            add_par(tf, it["text"], size=12, align=PP_ALIGN.CENTER, first=True,
                    color=GRAY, line=1.2)

def render_badge_grid(slide, z, l, t, w, h):
    """徽章网格：圆徽+标题+解释，每格独立圆角卡"""
    items = z["items"]; cols = z.get("cols", 3)
    rows = math.ceil(len(items) / cols)
    gap = 0.16
    cw = (w - (cols - 1) * gap) / cols
    ch = (h - (rows - 1) * gap) / rows
    for k, it in enumerate(items):
        r, c = k // cols, k % cols
        x, y = l + c * (cw + gap), t + r * (ch + gap)
        rbox(slide, x, y, cw, ch, fill=WHITE, line_color=LINE, lw=1.0)
        box(slide, x, y, cw, 0.055, fill=BLUE)
        badge(slide, x + 0.12, y + 0.16, 0.36, it.get("badge", str(k + 1)))
        tf = text_box(slide, x + 0.58, y + 0.12, cw - 0.70, 0.42, MSO_ANCHOR.MIDDLE)
        add_par(tf, it["title"], size=12.5, first=True, color=BLUE_D, bold=True)
        tf2 = text_box(slide, x + 0.14, y + 0.58, cw - 0.26, ch - 0.68,
                       MSO_ANCHOR.TOP)
        add_par(tf2, it["text"], size=12, first=True, line=1.26)

def render_compare_rows(slide, z, l, t, w, h):
    """对象行：平行四边形标签 + 维度卡 + 箭头 + 措施卡"""
    heads = z.get("heads", ["激励对象", "赋分维度 / 考核指标", "配套激励与应用"])
    rows = z["rows"]
    label_w, mid_w, arrow_w = 1.55, 5.35, 0.42
    right_w = w - label_w - 0.15 - mid_w - arrow_w - 0.14
    # 表头
    for hx, hw, htxt in ((l, label_w, heads[0]),
                         (l + label_w + 0.15, mid_w, heads[1]),
                         (l + label_w + 0.15 + mid_w + arrow_w + 0.14, right_w, heads[2])):
        tf = text_box(slide, hx, t, hw, 0.30, MSO_ANCHOR.MIDDLE)
        add_par(tf, htxt, size=12.5, align=PP_ALIGN.CENTER, first=True,
                color=BLUE_D, bold=True)
    y = t + 0.34
    gap = 0.12
    rh = (h - 0.34 - (len(rows) - 1) * gap) / len(rows)
    for row in rows:
        pl = shape(slide, MSO_SHAPE.PARALLELOGRAM, l, y + (rh - 0.52) / 2,
                   label_w, 0.52, fill=BLUE, adj=0.28)
        shape_text(pl, row["label"], size=12.5, align=PP_ALIGN.CENTER,
                   m=(0.12, 0.06, 0.01, 0.01), color=WHITE, bold=True)
        mx = l + label_w + 0.15
        mid = rbox(slide, mx, y, mid_w, rh, fill=CARD, line_color=LINE, lw=0.75)
        tfm = text_box(slide, mx + 0.10, y + 0.03, mid_w - 0.20, rh - 0.06,
                       MSO_ANCHOR.MIDDLE)
        add_par(tfm, row["mid"], size=12, first=True, line=1.24)
        ax = mx + mid_w + 0.03
        shape(slide, MSO_SHAPE.RIGHT_ARROW, ax, y + (rh - 0.30) / 2, arrow_w - 0.06,
              0.30, fill=GOLD)
        rx = mx + mid_w + arrow_w + 0.14
        rt = rbox(slide, rx, y, right_w, rh, fill=WHITE, line_color=BLUE, lw=1.0)
        tfr = text_box(slide, rx + 0.10, y + 0.03, right_w - 0.20, rh - 0.06,
                       MSO_ANCHOR.MIDDLE)
        add_par(tfr, row["right"], size=12, first=True, line=1.24)
        y += rh + gap

def render_timeline(slide, z, l, t, w, h):
    """时间轴：节点圆(月)+主题+连接线，下方任务逐条独立卡"""
    phases = z["phases"]; n = len(phases)
    colors = [BLUE, GOLD, RED][:n]
    gap = 0.25
    cw = (w - (n - 1) * gap) / n
    d = 0.62
    # 连接线
    box(slide, l + cw / 2, t + d / 2 - 0.015, w - cw, 0.03, fill=BLUE_L)
    for i, ph in enumerate(phases):
        x = l + i * (cw + gap)
        col = colors[i]
        tcol = WHITE if col != GOLD else BLUE_D
        node = shape(slide, MSO_SHAPE.OVAL, x + cw / 2 - d / 2, t, d, d,
                     fill=col, line_color=WHITE, lw=2.0)
        tf = node.text_frame; tf.word_wrap = False
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); _set_run(r, ph["month"], 14, True, tcol)
        tft = text_box(slide, x, t + d + 0.06, cw, 0.34, MSO_ANCHOR.MIDDLE)
        add_par(tft, ph["theme"], size=13, align=PP_ALIGN.CENTER, first=True,
                color=(GOLD_D if col == GOLD else col), bold=True)
        # 任务卡
        tasks = ph["tasks"]
        area_t = t + 1.10
        ests = [est_lines(tk, cw - 0.62, 12) * _lh(12) + 0.16 for tk in tasks]
        igap = 0.10
        total = sum(ests) + igap * (len(tasks) - 1)
        extra = max(0.0, (t + h - area_t) - total)
        y = area_t
        for j, tk in enumerate(tasks):
            ih = ests[j] + extra / len(tasks)
            rbox(slide, x, y, cw, ih, fill=CARD, line_color=LINE, lw=0.75)
            box(slide, x, y, 0.05, ih, fill=col)
            bd = 0.28
            badge(slide, x + 0.11, y + (ih - bd) / 2, bd, str(j + 1), fill=col,
                  tcolor=tcol, size=12)
            tfx = text_box(slide, x + 0.48, y + 0.02, cw - 0.58, ih - 0.04,
                           MSO_ANCHOR.MIDDLE)
            add_par(tfx, tk, size=12, first=True, line=1.22)
            y += ih + igap

def render_chips(slide, z, l, t, w, h):
    y = t
    if z.get("label"):
        box(slide, l, y + 0.03, 0.055, 0.26, fill=GOLD)
        tf = text_box(slide, l + 0.14, y, w - 0.14, 0.32, MSO_ANCHOR.MIDDLE)
        add_par(tf, z["label"], size=13, first=True, color=BLUE_D, bold=True)
        y += 0.36
    cx, row_h = l, 0.34
    for it in z["items"]:
        iw = _disp_len(it) * 0.165 + 0.78
        if cx + iw > l + w + 0.01:
            cx = l; y += row_h + 0.08
        chip = shape(slide, MSO_SHAPE.HEXAGON, cx, y, iw, row_h, fill=CHIP,
                     line_color=BLUE, lw=0.75)
        shape_text(chip, it, size=12, align=PP_ALIGN.CENTER,
                   m=(0.14, 0.14, 0.01, 0.01), color=BLUE_D)
        cx += iw + 0.14
    y += row_h + 0.08
    if z.get("text"):
        tf = text_box(slide, l, y, w, t + h - y, MSO_ANCHOR.TOP)
        add_par(tf, z["text"], size=12, first=True, line=1.26)

def _cell_text(cell, text, size, bold, color, align):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    tf = cell.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.2
        for seg_text, sb, sc in _segments(ln):
            r = p.add_run()
            _set_run(r, seg_text, size, bold or sb,
                     sc if sc is not None else color)

def render_table(slide, z, l, t, w, h):
    head, rows = z["head"], z["rows"]
    widths = z.get("widths") or [1] * len(head)
    tot = float(sum(widths))
    gf = slide.shapes.add_table(len(rows) + 1, len(head),
                                Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = gf.table
    tbl.first_row = False; tbl.horz_banding = False
    for ci, fr in enumerate(widths):
        tbl.columns[ci].width = Emu(int(w * fr / tot * 914400))
    tbl.rows[0].height = Inches(0.42)
    for ci, htxt in enumerate(head):
        c = tbl.cell(0, ci)
        c.fill.solid(); c.fill.fore_color.rgb = BLUE
        _cell_text(c, htxt, 12.5, True, WHITE, PP_ALIGN.CENTER)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri + 1, ci)
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if ri % 2 == 0 else CARD
            if ci == 0:
                _cell_text(c, val, 12, True, BLUE_D, PP_ALIGN.CENTER)
            else:
                _cell_text(c, val, 12, False, INK, PP_ALIGN.LEFT)

def render_banner(slide, z, l, t, w, h):
    b = box(slide, l, t, w, h, fill=BLUE)
    shape(slide, MSO_SHAPE.PARALLELOGRAM, l + w - 1.5, t, 1.5, h, fill=BLUE_D)
    shape(slide, MSO_SHAPE.PARALLELOGRAM, l + w - 1.0, t, 1.0, h, fill=GOLD, adj=0.35)
    box(slide, l, t, 0.10, h, fill=GOLD)
    tf = b.text_frame
    tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.28); tf.margin_right = Inches(1.8)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); _set_run(r, z["title"], 14, True, WHITE)
    for ln in z.get("lines", []):
        p = tf.add_paragraph(); p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(3)
        r = p.add_run(); _set_run(r, _plain(ln), 12, False, RGBColor(0xDC, 0xE6, 0xF5))

def render_note(slide, z, l, t, w, h):
    b = box(slide, l, t, w, h, fill=CARD)
    box(slide, l, t, 0.07, h, fill=RED)
    shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, l + 0.16, t + (h - 0.16) / 2,
          0.16, 0.16, fill=RED, adj=None).rotation = 90
    shape_text(b, z["text"], size=12.5, anchor=MSO_ANCHOR.MIDDLE,
               m=(0.42, 0.10, 0.03, 0.03))

RENDER = {"lead": render_lead, "kv_rows": render_kv_rows, "cards": render_cards,
          "flow_chain": render_flow_chain, "shape_row": render_shape_row,
          "badge_grid": render_badge_grid, "compare_rows": render_compare_rows,
          "timeline": render_timeline, "chips": render_chips,
          "table": render_table, "banner": render_banner, "note": render_note}

FLEX = {"kv_rows", "cards", "table", "flow_chain", "badge_grid",
        "compare_rows", "timeline", "shape_row"}

def render_stack(slide, zones, top=BODY_T, bottom=BODY_B):
    avail = bottom - top
    gap = 0.14
    ests = [est_zone(z) for z in zones]
    total = sum(ests) + gap * (len(zones) - 1)
    if total < avail:
        extra = avail - total
        flex_idx = [i for i, z in enumerate(zones) if z["type"] in FLEX]
        if flex_idx:
            base = sum(ests[i] for i in flex_idx) or 1
            for i in flex_idx:
                ests[i] += extra * ests[i] / base
    else:
        scale = (avail - gap * (len(zones) - 1)) / max(0.1, total - gap * (len(zones) - 1))
        ests = [e * scale for e in ests]
    y = top
    for z, h in zip(zones, ests):
        RENDER[z["type"]](slide, z, PAGE_L, y, PAGE_W, h)
        y += h + gap

# ---------------- 模板页克隆 ----------------

def _replace_deep(slide, mapping):
    cnt = 0
    def proc(sh):
        nonlocal cnt
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in sh.shapes: proc(sub)
            return
        if not sh.has_text_frame: return
        for pa in sh.text_frame.paragraphs:
            runs = list(pa.runs)
            if not runs: continue
            full = "".join(r.text for r in runs); new = full
            for o, v in mapping.items():
                if o in new: new = new.replace(o, str(v)); cnt += 1
            if new != full:
                runs[0].text = new
                for r in runs[1:]: r.text = ""
    for sh in slide.shapes: proc(sh)
    return cnt

def _clone(prs, src):
    n = prs.slides.add_slide(src.slide_layout)
    for ph in list(n.placeholders):
        n.shapes._spTree.remove(ph._element)
    for sh in src.shapes:
        n.shapes._spTree.append(deepcopy(sh._element))
    return n

# ---------------- 主流程 ----------------

def build(spec, output):
    prs = Presentation(TEMPLATE)
    tpl = list(prs.slides)

    cover = _clone(prs, tpl[0])
    _replace_deep(cover, {"此处输入一级标题": spec["title"],
                          "演讲人姓名": spec["speaker"]})
    chapters = spec.get("chapters", [])
    toc = _clone(prs, tpl[1])
    mapping = {f"输入{CH_NUM[i]}章节标题": ch["name"]
               for i, ch in enumerate(chapters[:5])}
    _replace_deep(toc, mapping)
    if len(chapters) < 5:
        extras = []
        for i in range(len(chapters), 5):
            extras += [f"输入{CH_NUM[i]}章节标题", CH_NUM[i]]
        sp = toc.shapes._spTree
        for sh in list(toc.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip() in extras:
                sp.remove(sh._element)

    for ci, ch in enumerate(chapters):
        div = _clone(prs, tpl[2])
        _replace_deep(div, {"输入一章节标题": f"{CH_NUM[ci]}、{ch['name']}"})
        for page in ch.get("pages", []):
            s = _clone(prs, tpl[3])
            _replace_deep(s, {"页面标题": page["title"]})
            render_stack(s, page["zones"])
            print(f"  OK {page['title'][:24]}")

    if spec.get("closing"):
        s = _clone(prs, tpl[3])
        for sh in list(s.shapes):
            if sh.has_text_frame and "页面标题" in sh.text_frame.text:
                s.shapes._spTree.remove(sh._element)
        render_closing(s, spec["closing"])
        print("  OK 封底")

    for sid in list(prs.slides._sldIdLst)[:5]:
        prs.slides._sldIdLst.remove(sid)
        prs.part.drop_rel(sid.rId)
    prs.save(output)
    print(f"OK {output} ({len(prs.slides._sldIdLst)}p)")

def render_closing(slide, page):
    box(slide, 0, 2.9, 13.333, 0.06, fill=GOLD)
    tf = text_box(slide, 1.0, 3.2, 11.33, 1.6, MSO_ANCHOR.TOP)
    p = add_par(tf, page.get("line1", "以上汇报，敬请批评指正"), size=32,
                align=PP_ALIGN.CENTER, first=True, color=BLUE_D, bold=True)
    if page.get("line2"):
        add_par(tf, page["line2"], size=14, align=PP_ALIGN.CENTER, color=GRAY)

if __name__ == "__main__":
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    out = sys.argv[2] if len(sys.argv) > 2 else "output.pptx"
    build(spec, out)
